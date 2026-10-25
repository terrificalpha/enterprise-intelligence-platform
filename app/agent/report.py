"""Report generation tool: builds a Markdown + DOCX report from analytics, ML, and cited evidence."""
from __future__ import annotations

import re
import time

from docx import Document

from ..config import settings
from ..llm.generator import readable


def build_report(tb, title: str, customer_id: str | None, top_n: int) -> dict:
    settings.ensure_dirs()
    ov = tb.tool_analyze_data("overview")
    by_ind = tb.tool_analyze_data("churn_by", column="industry")
    rev = tb.tool_analyze_data("revenue_at_risk", column="industry")
    top = tb.tool_analyze_data("top_risk", n=top_n)
    cor = tb.tool_analyze_data("correlations")
    metrics = tb.pred.metrics()
    sel = metrics.get("selected", {})
    cites: list = []

    def cite(hit) -> int:
        for i, h in enumerate(cites, 1):
            if h.chunk.chunk_id == hit.chunk.chunk_id:
                return i
        cites.append(hit)
        return len(cites)

    md = [f"# {title}", f"*Generated {time.strftime('%Y-%m-%d %H:%M')} by the Enterprise Intelligence Platform*", ""]
    md += ["## 1. Executive summary", ov["summary"], by_ind["summary"], rev["summary"], ""]
    md += ["## 2. Portfolio KPIs", "```", *[f"{k}: {v}" for k, v in ov["result"].items()], "```", ""]
    md += ["## 3. Churn by industry", by_ind["markdown"], "", "## 4. Revenue at risk by industry", rev["markdown"], ""]
    md += [f"## 5. Top {top_n} at-risk accounts", top["markdown"], "", "## 6. Churn drivers (data)", cor["summary"], "", cor["markdown"], ""]
    # policy-grounded actions for the portfolio
    md += ["## 7. Policy-grounded recommendations"]
    for q in ["High Risk accounts executive outreach save plan", "retention discount approval authority",
              "executive sponsor accounts annual contract value", "Credit Watch late payments joint review"]:
        hits = tb.kb.search(q, k=3)
        pol = next((h for h in hits if h.chunk.doc_type in ("pdf", "docx") and h.coverage >= 0.4), None)
        if pol:
            md.append(f"- {readable(pol)} [{cite(pol)}]")
    md.append("")
    if customer_id:
        b = tb.ins.briefing(customer_id)
        if "error" not in b:
            md += [f"## 8. Account deep-dive: {b['customer_id']}", ""]
            offset = len(cites)
            for c in b["citations"]:
                class _H:  # adapt briefing citations into the report's reference list
                    chunk = type("C", (), {"chunk_id": f"b{c['id']}", "source": c["source"], "location": c["location"]})()
                    text = c["snippet"]
                cites.append(_H)
            narr = re.sub(r"\[(\d+)\]", lambda m: f"[{int(m.group(1)) + offset}]", b["narrative"])
            md += [narr, ""]
    md += ["## Model & methodology",
           f"- Selected model: **{sel.get('model', 'n/a')}** (calibrated). Test ROC-AUC {sel.get('test_calibrated', {}).get('roc_auc', float('nan')):.3f}, "
           f"F1 {sel.get('test_calibrated', {}).get('f1', float('nan')):.3f}.",
           "- Churn probabilities are calibrated; risk tiers follow Customer Retention Policy s.2 (High >= 0.60, Medium >= 0.30).",
           "- Evidence retrieved with hybrid search (dense + BM25, reciprocal rank fusion) and reranking.", "",
           "## Limitations",
           "- Data is synthetic; probabilities describe model risk, not certainty. Policy citations are lexically verified, not semantically proven.", "",
           "## References"]
    for i, h in enumerate(cites, 1):
        txt = getattr(h, "text", None) or readable(h)
        md.append(f"[{i}] {h.chunk.source} ({h.chunk.location}): {txt[:200]}")
    text = "\n".join(md)
    slug = re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")[:40] or "report"
    stamp = time.strftime("%Y%m%d-%H%M%S")
    md_path = settings.reports_dir / f"{stamp}_{slug}.md"
    md_path.write_text(text)
    docx_path = md_path.with_suffix(".docx")
    _to_docx(text, docx_path)
    return {"markdown_path": str(md_path), "docx_path": str(docx_path), "preview": text[:1500],
            "summary": f"Report written to {md_path.name} ({len(text):,} chars, {len(cites)} references)."}


def _to_docx(md: str, path) -> None:
    d = Document()
    in_code = False
    table_rows: list[list[str]] = []

    def flush_table():
        nonlocal table_rows
        if table_rows:
            t = d.add_table(rows=0, cols=len(table_rows[0]))
            t.style = "Light Grid Accent 1"
            for r in table_rows:
                cells = t.add_row().cells
                for c, v in zip(cells, r):
                    c.text = v
            table_rows = []

    for line in md.splitlines():
        if line.startswith("```"):
            in_code = not in_code
            continue
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if not set("".join(cells)) <= set("-: "):
                table_rows.append(cells)
            continue
        flush_table()
        if line.startswith("# "):
            d.add_heading(line[2:], 0)
        elif line.startswith("## "):
            d.add_heading(line[3:], 1)
        elif re.match(r"^\d+\. ", line):
            d.add_paragraph(re.sub(r"\*\*|\*", "", line), style="List Number")
        elif line.startswith("- "):
            d.add_paragraph(re.sub(r"\*\*|\*", "", line[2:]), style="List Bullet")
        elif line.strip():
            d.add_paragraph(re.sub(r"\*\*|\*", "", line))
    flush_table()
    d.save(path)
