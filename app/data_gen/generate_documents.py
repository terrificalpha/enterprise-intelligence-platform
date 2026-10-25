"""Generate the enterprise document corpus (PDF, DOCX, CSV) with known, checkable facts.

Facts in the policies are the ground truth used by the RAG evaluation set (data/eval/*.json).
The Q3 review report is computed from the real customer data so documents and data agree.
"""
from __future__ import annotations

import numpy as np
import pandas as pd
from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..config import settings

# --------------------------------------------------------------------------- PDF policies
RETENTION = ("Customer Retention Policy", [
    ("1. Purpose and Scope", [
        "This policy defines how the Customer Success organization identifies, escalates and mitigates churn risk "
        "across all B2B accounts. It applies to Account Managers, Customer Success Managers and Regional VPs."]),
    ("2. Risk Tiers", [
        "Accounts are classified using the churn-risk score produced by the Enterprise Intelligence Platform. "
        "High Risk: predicted churn probability of 0.60 or higher. Medium Risk: predicted churn probability from 0.30 "
        "up to 0.60. Low Risk: predicted churn probability below 0.30."]),
    ("3. Escalation Requirements", [
        "High Risk accounts must receive an executive outreach call within 5 business days. A written save plan must "
        "be approved by the Regional VP within 10 business days. Medium Risk accounts require a customer success "
        "manager check-in within 15 business days."]),
    ("4. Discount Authority", [
        "Account Managers may approve retention discounts of up to 10 percent of annual contract value. Discounts "
        "between 10 and 20 percent require Regional VP approval. Discounts above 20 percent require CFO approval. "
        "No retention discount may exceed 30 percent under any circumstances."]),
    ("5. Executive Sponsor Program", [
        "Accounts with annual contract value of 100,000 USD or more must be assigned an executive sponsor. "
        "Executive sponsors hold a quarterly business review with the customer."]),
    ("6. Renewal Management", [
        "Renewal conversations must begin 120 days before the contract end date. Accounts within 60 days of renewal "
        "with declining usage are flagged as Renewal Watch and reviewed weekly."]),
    ("7. Adoption Recovery", [
        "When seat utilization falls below 50 percent, a training and enablement plan must be offered within 30 days. "
        "A usage decline of more than 20 percent over 90 days triggers a customer health review."]),
])

CREDIT = ("Credit and Payment Risk Policy", [
    ("1. Payment Terms", [
        "Standard payment terms are Net 30. Net 45 terms are available only to Enterprise customers with CFO "
        "approval. Terms longer than Net 45 are not permitted."]),
    ("2. Late Payment Monitoring", [
        "Three or more late payments within 12 months place the account on Credit Watch. Five or more late payments "
        "trigger a formal credit review and may require prepayment for renewals."]),
    ("3. Collections Process", [
        "Invoices 30 days overdue receive an automated reminder. Invoices 45 days overdue are escalated to the "
        "Collections team. Invoices 60 days overdue trigger a credit hold on new orders and service expansion."]),
    ("4. Late Fees", [
        "A late fee of 1.5 percent per month is charged on overdue balances, unless waived by the Finance Director."]),
    ("5. Interaction with Retention", [
        "Credit Watch accounts with high churn risk require a joint review by Finance and Customer Success before "
        "any retention discount is offered."]),
])

SUPPORT = ("Support Service Level Policy", [
    ("1. Priority Definitions and Targets", [
        "P1 Critical: first response within 1 hour and resolution within 8 hours. P2 High: first response within "
        "4 hours and resolution within 24 hours. P3 Normal: first response within 8 business hours and resolution "
        "within 72 hours. P4 Low: first response within 2 business days and resolution within 10 business days."]),
    ("2. SLA Breach Credits", [
        "Each SLA breach earns the customer a service credit of 5 percent of the monthly fee, capped at 20 percent "
        "of the monthly fee."]),
    ("3. Escalation", [
        "Three or more SLA breaches in a rolling 90-day window trigger an executive escalation and a root cause "
        "analysis delivered within 7 days."]),
    ("4. Quality Targets", [
        "The target average ticket resolution time is below 24 hours. Customers with an average resolution time "
        "above 24 hours are reviewed monthly by the Head of Support."]),
])

# --------------------------------------------------------------------------- DOCX
MSA = ("Master Service Agreement - Standard Terms", [
    ("1. Term and Renewal", [
        "The initial term is 12 months. The agreement renews automatically for successive 12-month terms unless "
        "either party gives written notice of non-renewal at least 60 days before the end of the current term."]),
    ("2. Termination", [
        "Either party may terminate for convenience with 90 days written notice, except for Multi-year customers who "
        "are committed for the full term. Either party may terminate for cause if a material breach is not cured "
        "within 30 days of notice."]),
    ("3. Fees and Price Adjustments", [
        "Fees are subject to an annual price increase of 5 percent applied at each renewal. Invoices are payable "
        "Net 30 from the invoice date."]),
    ("4. Limitation of Liability", [
        "Each party's aggregate liability is capped at the fees paid in the preceding 12 months."]),
    ("5. Data Protection and Retention", [
        "Customer data is deleted within 30 days after termination. Personal data is processed under the Data "
        "Processing Addendum. The provider must notify the customer of a personal data breach within 72 hours."]),
    ("6. Service Levels", [
        "The service commits to 99.9 percent monthly uptime. If uptime falls below 99.9 percent a 10 percent service "
        "credit applies; below 99.0 percent a 25 percent credit applies."]),
])

COMPETITORS = ("Competitive Landscape Brief", [
    ("1. Overview", [
        "Three competitors are most frequently mentioned in lost and at-risk deals: NimbusOps, Vantiq Cloud and "
        "OrbitDesk."]),
    ("2. NimbusOps", [
        "NimbusOps undercuts our mid-market pricing by 15 to 20 percent and is strongest in the Retail segment. "
        "Their weakness is limited enterprise integrations."]),
    ("3. Vantiq Cloud", [
        "Vantiq Cloud focuses on the Manufacturing industry and offers on-premise deployment, which appeals to "
        "customers with strict data residency needs."]),
    ("4. OrbitDesk", [
        "OrbitDesk bundles free premium support for the first year of any contract, which has attracted customers "
        "frustrated by slow resolution times."]),
    ("5. Recommended Defence", [
        "Defend at-risk accounts with value-based selling and multi-year terms rather than deep discounts. Highlight "
        "integration depth, the executive sponsor program and our SLA credits."]),
])


def _write_pdf(path, spec) -> None:
    title, sections = spec
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=10.5, leading=15)
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm,
                            topMargin=2.2 * cm, bottomMargin=2.2 * cm, title=title, author="Enterprise Policy Office")
    story = [Paragraph(title, styles["Title"]), Paragraph("Version 3.2 - Effective 1 January 2026 - Internal", styles["Italic"]),
             Spacer(1, 14)]
    for heading, paras in sections:
        story.append(Paragraph(heading, styles["Heading2"]))
        for p in paras:
            story.append(Paragraph(p, body))
            story.append(Spacer(1, 6))
    doc.build(story)


def _write_docx(path, spec, extra=None) -> None:
    title, sections = spec
    d = Document()
    d.add_heading(title, level=0)
    for heading, paras in sections:
        d.add_heading(heading, level=1)
        for p in paras:
            d.add_paragraph(p)
    if extra:
        extra(d)
    d.save(path)


def _q3_report(path, stats: dict) -> None:
    ind = stats["by_industry"]
    top2 = list(ind.sort_values(ascending=False).index[:2])
    spec = ("Q3 Account Review Report", [
        ("1. Executive Summary", [
            f"The overall customer churn rate in the portfolio is {stats['churn']:.1%}. Across {stats['n']:,} accounts, "
            f"{top2[0]} and {top2[1]} show the highest churn at {ind[top2[0]]:.1%} and {ind[top2[1]]:.1%} respectively. "
            f"Annual contract value of churned accounts totals {stats['lost_acv']/1e6:,.1f} million USD."]),
        ("2. Key Churn Drivers", [
            "Analysis of churned accounts shows four recurring drivers: repeated SLA breaches, three or more late "
            "payments in twelve months, declining product usage, and active competitor mentions. Accounts on Monthly "
            "contracts churn far more often than Multi-year accounts."]),
        ("3. Recommendations from Leadership Review", [
            "Leadership asked Customer Success to prioritise High Risk accounts by revenue at risk, assign executive "
            "sponsors to all accounts above 100,000 USD, and coordinate with Finance on Credit Watch accounts before "
            "any discount is offered."]),
    ])

    def extra(d):
        d.add_heading("4. Churn Rate by Industry", level=1)
        t = d.add_table(rows=1, cols=2)
        t.style = "Light Grid Accent 1"
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Industry", "Churn rate"
        for k, v in ind.sort_values(ascending=False).items():
            r = t.add_row().cells
            r[0].text, r[1].text = str(k), f"{v:.1%}"
    _write_docx(path, spec, extra)


# --------------------------------------------------------------------------- CSV (tickets, notes)
_TICKET_TEMPLATES = {
    "slow": ("Slow ticket response", "Customer reports that ticket {n} has been open for several days without a resolution. Escalation requested."),
    "outage": ("Service outage", "Production dashboard was unavailable for {h} hours. Customer requests a root cause analysis and SLA credit."),
    "billing": ("Invoice dispute", "Customer disputes the latest invoice and asks for clarification of the price increase applied at renewal."),
    "feature": ("Missing feature request", "Customer asks for a deeper ERP integration and says competitors already provide it."),
    "training": ("Training request", "New team members need onboarding sessions; seat adoption is lower than expected."),
    "perf": ("Performance degradation", "Reports take over a minute to load during peak hours. Customer is frustrated."),
}
_NOTE_RISK = [
    "Champion left the company; new contact is evaluating alternatives including {c}.",
    "Customer mentioned budget cuts and asked about downgrading seats at renewal.",
    "{c} presented a proposal priced well below our renewal quote; customer asked us to match it.",
    "Frustrated by repeated SLA breaches; asked for an executive call before renewal.",
    "Payments have been late three quarters in a row; finance contact cites cash-flow pressure.",
    "Usage has dropped sharply since the reorganisation; only a fraction of seats are active.",
]
_NOTE_OK = [
    "Quarterly business review went well; customer plans to expand to a second region.",
    "Customer is very satisfied with the new reporting module and offered to be a reference.",
    "Executive sponsor met the CIO; renewal expected to proceed on current terms.",
    "Training workshop completed; adoption metrics improving week over week.",
]


def _tickets_and_notes(cust: pd.DataFrame, rng) -> tuple[pd.DataFrame, pd.DataFrame]:
    comp = ["NimbusOps", "Vantiq Cloud", "OrbitDesk"]
    t_rows, n_rows, tid = [], [], 1
    sample = cust.sample(180, random_state=7)
    for _, r in sample.iterrows():
        for _ in range(rng.integers(1, 3)):
            key = rng.choice(list(_TICKET_TEMPLATES))
            subj, desc = _TICKET_TEMPLATES[key]
            t_rows.append({"ticket_id": f"T{tid:05d}", "customer_id": r.customer_id,
                           "priority": rng.choice(["P1", "P2", "P3", "P4"], p=[0.1, 0.3, 0.45, 0.15]),
                           "subject": subj, "description": desc.format(n=int(rng.integers(1000, 9999)), h=int(rng.integers(2, 12))),
                           "resolution_hours": round(float(rng.gamma(3, 8)), 1)})
            tid += 1
    n_sample = cust.sample(200, random_state=11)
    for _, r in n_sample.iterrows():
        risky = r.churned == 1 or rng.random() < 0.15
        txt = rng.choice(_NOTE_RISK if risky else _NOTE_OK).format(c=rng.choice(comp))
        n_rows.append({"note_id": f"N{len(n_rows)+1:05d}", "customer_id": r.customer_id,
                       "date": f"2026-{int(rng.integers(1, 10)):02d}-{int(rng.integers(1, 28)):02d}",
                       "author": rng.choice(["A. Rao", "M. Chen", "S. Okafor", "L. Meyer"]), "note": txt})
    return pd.DataFrame(t_rows), pd.DataFrame(n_rows)


def main() -> None:
    settings.ensure_dirs()
    rng = np.random.default_rng(settings.random_state)
    out = settings.documents_dir
    cust = pd.read_csv(settings.customers_csv).drop_duplicates("customer_id")
    from ..ml.features import normalise_category
    cust["industry"] = normalise_category(cust["industry"])
    _write_pdf(out / "Customer_Retention_Policy.pdf", RETENTION)
    _write_pdf(out / "Credit_and_Payment_Risk_Policy.pdf", CREDIT)
    _write_pdf(out / "Support_Service_Level_Policy.pdf", SUPPORT)
    _write_docx(out / "Master_Service_Agreement.docx", MSA)
    _write_docx(out / "Competitive_Landscape_Brief.docx", COMPETITORS)
    stats = {"n": len(cust), "churn": cust.churned.mean(),
             "by_industry": cust.groupby("industry").churned.mean(),
             "lost_acv": cust.loc[cust.churned == 1, "annual_contract_value"].sum()}
    _q3_report(out / "Q3_Account_Review_Report.docx", stats)
    tickets, notes = _tickets_and_notes(cust, rng)
    tickets.to_csv(out / "support_tickets.csv", index=False)
    notes.to_csv(out / "account_notes.csv", index=False)
    print("Documents written to", out)
    for p in sorted(out.iterdir()):
        print("  ", p.name)


if __name__ == "__main__":
    main()
