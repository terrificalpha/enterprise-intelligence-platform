import tempfile
import unittest
from pathlib import Path

from docx import Document
from reportlab.platypus import Paragraph, SimpleDocTemplate
from reportlab.lib.styles import getSampleStyleSheet

from tests.helpers import ensure_artifacts


class TestIngestion(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        ensure_artifacts()
        cls.tmp = Path(tempfile.mkdtemp())

    def test_pdf_extraction_keeps_sections_and_pages(self):
        from app.ingestion.loaders import load_file
        from app.config import settings
        ch = load_file(settings.documents_dir / "Customer_Retention_Policy.pdf")
        self.assertGreaterEqual(len(ch), 7)
        risk = [c for c in ch if "Risk Tiers" in c.section]
        self.assertTrue(risk and "0.60" in risk[0].text)
        self.assertTrue(all(c.source == "Customer_Retention_Policy.pdf" and c.location.startswith("page") for c in ch))

    def test_docx_extraction_with_headings_and_tables(self):
        from app.ingestion.loaders import load_file
        d = Document(); d.add_heading("Plan", 0); d.add_heading("1. Budget", 1); d.add_paragraph("The budget is 500 dollars.")
        t = d.add_table(rows=2, cols=2); t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Team", "Cost"
        t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Alpha", "42"
        p = self.tmp / "plan.docx"; d.save(p)
        ch = load_file(p)
        txt = " ".join(c.text for c in ch)
        self.assertIn("budget is 500 dollars", txt)
        self.assertIn("Team: Alpha", txt)            # table row flattened with its header
        self.assertTrue(any(c.section.startswith("1. Budget") for c in ch))

    def test_csv_rows_become_records_with_customer_id(self):
        from app.ingestion.loaders import load_file
        p = self.tmp / "n.csv"
        p.write_text("customer_id,note\nC0001,Budget cuts planned\nC0002,Happy customer\n")
        ch = load_file(p)
        self.assertEqual(len(ch), 2)
        self.assertEqual(ch[0].customer_id, "C0001")
        self.assertIn("Budget cuts", ch[0].text)
        self.assertEqual(ch[1].location, "row 2")

    def test_chunking_respects_size_and_overlap(self):
        from app.ingestion.loaders import split_text
        text = " ".join(f"Sentence number {i} is here." for i in range(60))
        parts = split_text(text, 300, 60)
        self.assertGreater(len(parts), 3)
        self.assertTrue(all(len(p) <= 420 for p in parts))
        self.assertEqual(" ".join(parts).count("Sentence number 59"), 1 + (1 if parts[-2].endswith("59 is here.") else 0))

    def test_unsupported_type_rejected(self):
        from app.ingestion.loaders import load_file
        p = self.tmp / "x.exe"; p.write_bytes(b"x")
        with self.assertRaises(ValueError):
            load_file(p)

    def test_corpus_covers_all_three_formats(self):
        from app.ingestion.loaders import load_directory
        types = {c.doc_type for c in load_directory()}
        self.assertTrue({"pdf", "docx", "csv"} <= types)


if __name__ == "__main__":
    unittest.main()
